#!/usr/bin/env python3
"""Classe des entreprises (export Kompass) selon leur activité IA / Data, d'après leur site web.

Étapes :
  1. crawl  : page d'accueil + jusqu'à 3 pages internes pertinentes par site, texte mis en cache (SQLite)
  2. score  : score IA et score Data à partir de mots-clés FR/EN pondérés (recalculable sans re-télécharger)
  3. export : nouveau fichier Excel = colonnes d'origine + résultats (le fichier source n'est jamais modifié)

Exemples :
  python classify_ai_data.py data/potentiel.xlsx --sample 20          # test sur 20 entreprises variées
  python classify_ai_data.py data/potentiel.xlsx                      # tout le fichier (reprend le cache)
  python classify_ai_data.py data/potentiel.xlsx --rescore-only       # ré-applique les mots-clés au cache
"""
from __future__ import annotations

import argparse
import asyncio
import json
import re
import sqlite3
import sys
import time
import unicodedata
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urldefrag, urljoin, urlparse
from urllib.robotparser import RobotFileParser

import httpx
import pandas as pd
from bs4 import BeautifulSoup
from tqdm import tqdm

# --------------------------------------------------------------------------- réglages réseau
USER_AGENT = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"
)
HEADERS = {
    "User-Agent": USER_AGENT,
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
    "Accept-Language": "fr-FR,fr;q=0.9,en;q=0.8",
}
TIMEOUT = 15.0
MAX_PARALLEL_REQUESTS = 10
MIN_DELAY_PER_DOMAIN = 1.0  # secondes entre deux requêtes vers un même hôte
MAX_REDIRECTS = 10
MAX_BYTES = 3_000_000
MAX_INTERNAL_PAGES = 3
MAX_TEXT_CHARS = 150_000  # par page, pour garder un cache raisonnable
MIN_WORDS = 40  # en dessous : site considéré "sans texte" (souvent un site 100 % JavaScript)

# --------------------------------------------------------------------------- liens internes
# Pages à ignorer : elles parlent de "données" sans rapport avec le métier.
EXCLUDE_LINK = re.compile(
    r"mentions?[-_ ]?l[eé]gales?|legal|cgu|cgv|conditions[-_ ]g[eé]n[eé]rales|terms|"
    r"cookie|confidentialit|privacy|rgpd|gdpr|donnees[-_]personnelles|protection[-_]des[-_]donnees|"
    r"recrut|carri[eè]re|career|jobs?\b|emploi|offres[-_]d[-_]emploi|candidat|rejoign|join[-_]?us|"
    r"nous[-_]rejoindre|talent|stage|alternance|"
    r"contact|login|connexion|sign[-_]?in|account|panier|cart|plan[-_]du[-_]site|sitemap|"
    r"accessibilit|press|presse|actualit|news|blog|evenement|event|webinar",
    re.I,
)
SKIP_EXT = re.compile(r"\.(pdf|jpe?g|png|gif|svg|webp|zip|docx?|xlsx?|pptx?|mp4|mp3|avi|css|js|xml|ico)$", re.I)
# Pages souhaitées, par ordre de priorité décroissante (poids).
WANTED_LINK = [
    (re.compile(r"expertise|savoir[-_ ]?faire|comp[eé]tence|m[eé]tier", re.I), 6),
    (re.compile(r"services?|prestations?|offres?|what[-_ ]we[-_ ]do", re.I), 6),
    (re.compile(r"solutions?|produits?|products?|platform|plateforme", re.I), 5),
    (re.compile(r"\b(data|ia|ai|intelligence)\b|data[-_]|[-_/]ia\b|[-_/]ai\b", re.I), 5),
    (re.compile(r"domaines?|activit[eé]s?|secteurs?|industr|use[-_ ]?cases?|cas[-_ ]d", re.I), 3),
    (re.compile(r"[aà][-_ ]propos|about|qui[-_ ]sommes|notre[-_ ](soci[eé]t[eé]|entreprise|histoire)|"
                r"entreprise|company|who[-_ ]we[-_ ]are|cabinet|groupe", re.I), 4),
]

# Éléments HTML à retirer avant l'extraction du texte.
CONSENT_PATTERN = re.compile(
    r"cookie|consent|gdpr|rgpd|tarteaucitron|didomi|onetrust|axeptio|cookiebot|cmplz|"
    r"usercentrics|quantcast|klaro|iubenda|termly|cc[-_]banner|cc[-_]window",
    re.I,
)

# --------------------------------------------------------------------------- mots-clés
# (motif regex, libellé, poids). Les motifs sont appliqués à un texte en minuscules sans accents,
# sauf ceux marqués CS (sensibles à la casse) appliqués au texte original : ils servent pour
# les sigles courts ambigus ("IA", "AI", "BI") qu'on ne veut pas confondre avec des mots courants.
STRONG, MEDIUM, WEAK = 4.0, 2.0, 0.5


@dataclass(frozen=True)
class KW:
    pattern: str
    label: str
    weight: float
    case_sensitive: bool = False


IA_KEYWORDS = [
    KW(r"intelligences? artificielles?", "intelligence artificielle", STRONG),
    KW(r"artificial intelligence", "artificial intelligence", STRONG),
    KW(r"machine[- ]learning", "machine learning", STRONG),
    KW(r"apprentissage (automatique|machine)", "apprentissage automatique", STRONG),
    KW(r"deep[- ]learning", "deep learning", STRONG),
    KW(r"apprentissage profond", "apprentissage profond", STRONG),
    KW(r"r[eé]seaux? (de )?neurones|neural networks?", "réseaux de neurones", STRONG),
    KW(r"computer vision", "computer vision", STRONG),
    KW(r"vision (par ordinateur|artificielle)", "vision par ordinateur", STRONG),
    KW(r"traitement (automatique )?du langage( naturel)?", "traitement du langage naturel", STRONG),
    KW(r"natural language (processing|understanding)", "natural language processing", STRONG),
    KW(r"\bNLP\b", "NLP", STRONG, True),
    KW(r"\bLLMs?\b", "LLM", STRONG, True),
    KW(r"large language models?|grands? modeles? de langage", "large language model", STRONG),
    KW(r"\bia generatives?|intelligence artificielle generative", "IA générative", STRONG),
    KW(r"generative ai|\bgen ?ai\b", "generative AI", STRONG),
    KW(r"\bmlops\b|\bllmops\b", "MLOps", STRONG),
    KW(r"\bagents? (ia|ai|intelligents?|autonomes?)\b|\bagentic\b|\bia agentique", "agents IA", STRONG),
    KW(r"retrieval[- ]augmented generation", "RAG", STRONG),
    KW(r"\b(chat ?gpt|gpt-?[345]o?|openai|mistral ai|hugging ?face|langchain|tensorflow|pytorch|scikit-learn)\b",
       "outils IA (GPT, PyTorch...)", STRONG),
    KW(r"\bIA\b", "IA", MEDIUM, True),
    KW(r"\bA\.?I\.?\b(?![-\.]?\w)", "AI", MEDIUM, True),
    KW(r"\bchat ?bots?\b|agents? conversationnels?|assistants? (virtuels?|conversationnels?)", "chatbot", MEDIUM),
    KW(r"\b(mod[eè]les?|analyses?|maintenance) pr[eé]dictive?s?|predictive (model|analytics|maintenance)s?",
       "prédictif", MEDIUM),
    KW(r"reconnaissance (d.images|vocale|faciale|de formes)|image recognition|speech recognition",
       "reconnaissance image/voix", MEDIUM),
    KW(r"\bai[- ](powered|driven|based|first)\b|\bpowered by ai\b|\bbasee? sur l.ia\b", "AI-powered", MEDIUM),
    KW(r"\bcopilot\b|\bprompt(s|ing)?\b", "copilot / prompt", MEDIUM),
    KW(r"\balgorithmes?\b|\balgorithms?\b", "algorithme", WEAK),
    KW(r"\bintelligent(e|s|es)?\b|\bsmart\b", "intelligent / smart", WEAK),
    KW(r"\bautomatisation\b|\bautomation\b", "automatisation", WEAK),
    KW(r"\binnovation\b|\binnovant", "innovation", WEAK),
]

DATA_KEYWORDS = [
    KW(r"data[- ]scien(ce|tists?)", "data science", STRONG),
    KW(r"science des donnees", "science des données", STRONG),
    KW(r"data[- ]engineer(ing|s)?|ingenieurs? (de |des )?donnees|ingenierie (de |des )?donnees", "data engineering",
       STRONG),
    KW(r"data[- ]analysts?|analystes? (de )?donnees", "data analyst", STRONG),
    KW(r"big[- ]data", "big data", STRONG),
    KW(r"business[- ]intelligence|informatique decisionnelle|\bdecisionnel(le)?\b", "business intelligence", STRONG),
    KW(r"data[- ]?(lake|lakehouse|warehouse|wharehouse|mart|hub|mesh|fabric|stack|platform|pipelines?|ops)\b",
       "data lake / warehouse / platform", STRONG),
    KW(r"entrepots? de donnees|lacs? de donnees", "entrepôt de données", STRONG),
    KW(r"data[- ]?(governance|management|quality|catalog)|gouvernance (de la |des )donnees|"
       r"qualite (de la |des )donnees|master data|\bmdm\b", "gouvernance / qualité data", STRONG),
    KW(r"data[- ]?vi[sz]|data[- ]?visuali[sz]ation|visualisation de(s)? donnees", "dataviz", STRONG),
    KW(r"\bpower ?bi\b|\bqlik(view| sense)?\b|\btableau (software|desktop|server|cloud)\b|\blooker\b|"
       r"\bmicrostrategy\b|\bbusiness ?objects\b|\bcognos\b", "outils BI (Power BI, Qlik...)", STRONG),
    KW(r"\bsnowflake\b|\bdatabricks\b|\bhadoop\b|\bapache spark\b|\bpyspark\b|\bkafka\b|\bbigquery\b|\btalend\b|"
       r"\binformatica\b|\bdataiku\b|\bairflow\b|\bdbt\b|\bredshift\b|\bsynapse\b|\bazure data factory\b|\bsas viya\b",
       "outils data (Snowflake, Databricks...)", STRONG),
    KW(r"\banalytics\b|data[- ]analytics|analyse (de|des) donnees|analyse de la donnee|advanced analytics",
       "analytics", STRONG),
    KW(r"data[- ]driven|pilotage par la donnee|culture (de la )?data|strategie data|valorisation (de vos |des )?donnees|"
       r"valoriser (vos|les) donnees|exploitation (de vos |des )donnees", "data-driven / valorisation", MEDIUM),
    KW(r"\bETL\b|\bELT\b", "ETL", MEDIUM, True),
    KW(r"\bBI\b", "BI", MEDIUM, True),
    KW(r"\bdata\b", "data", MEDIUM),
    KW(r"tableaux? de bord|dashboards?|\breporting\b|\bkpis?\b", "reporting / dashboard", WEAK),
    KW(r"\bdonnees\b", "données", WEAK),
    KW(r"\bdigital(e|es|isation)?\b|\bnumerique\b", "digital", WEAK),
    KW(r"\bstatisti(que|ques|cal|cs)\b", "statistiques", WEAK),
]

MAX_COUNT_PER_KEYWORD = 3  # au-delà, les répétitions (menu/pied de page sur chaque page) ne comptent plus
MAX_WEAK_TOTAL = 2.0       # plafond des termes faibles : à eux seuls ils ne suffisent jamais

# Seuils de décision (voir decide()).
THRESHOLD_YES = 14.0
THRESHOLD_PROBABLE = 6.0


def strip_accents(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def _compile(keywords: list[KW]):
    out = []
    for k in keywords:
        pat = k.pattern if k.case_sensitive else strip_accents(k.pattern)
        out.append((re.compile(pat, 0 if k.case_sensitive else re.I), k))
    return out


IA_COMPILED = _compile(IA_KEYWORDS)
DATA_COMPILED = _compile(DATA_KEYWORDS)


def score_text(text: str, compiled) -> dict:
    """Retourne score, nombre de termes forts distincts, et détail des mots-clés trouvés."""
    norm = strip_accents(text).lower()
    hits: dict[str, int] = defaultdict(int)
    weights: dict[str, float] = {}
    for rx, kw in compiled:
        n = len(rx.findall(text if kw.case_sensitive else norm))
        if n:
            hits[kw.label] += n
            weights[kw.label] = max(weights.get(kw.label, 0), kw.weight)
    strong = medium = weak = 0.0
    n_strong = 0
    for label, n in hits.items():
        w = weights[label]
        contrib = w * min(n, MAX_COUNT_PER_KEYWORD)
        if w >= STRONG:
            strong += contrib
            n_strong += 1
        elif w >= MEDIUM:
            medium += contrib
        else:
            weak += contrib
    score = strong + medium + min(weak, MAX_WEAK_TOTAL)
    detail = sorted(hits.items(), key=lambda kv: (-weights[kv[0]], -kv[1]))
    return {
        "score": round(score, 1),
        "n_strong": n_strong,
        "strong_score": strong,
        "keywords": "; ".join(f"{k}×{n}" if n > 1 else k for k, n in detail),
        "strong_keywords": "; ".join(k for k, _ in detail if weights[k] >= STRONG),
    }


def decide(s: dict) -> str:
    """oui : au moins 2 termes forts distincts et un score élevé ;
    probable : au moins un terme fort, ou beaucoup de termes moyens ;
    non : sinon (les termes faibles seuls ne suffisent jamais)."""
    if s["n_strong"] >= 2 and s["score"] >= THRESHOLD_YES:
        return "oui"
    if s["n_strong"] >= 1 and s["score"] >= THRESHOLD_PROBABLE:
        return "probable"
    if s["n_strong"] == 0 and s["score"] - min(s["score"], MAX_WEAK_TOTAL) >= 2 * THRESHOLD_PROBABLE:
        return "probable"
    return "non"


# --------------------------------------------------------------------------- extraction HTML
def extract_text(soup: BeautifulSoup) -> str:
    for tag in soup(["script", "style", "noscript", "svg", "iframe", "template", "canvas", "form", "button"]):
        tag.decompose()
    for tag in soup.find_all(True):
        if getattr(tag, "decomposed", False):
            continue
        attrs = getattr(tag, "attrs", None)
        if not attrs:
            continue
        ident = " ".join([str(attrs.get("id", ""))] + [str(c) for c in attrs.get("class", []) or []])
        if ident.strip() and CONSENT_PATTERN.search(ident):
            tag.decompose()
    parts = []
    if soup.title and soup.title.string:
        parts.append(soup.title.string)
    for name in ("description", "og:description", "keywords"):
        m = soup.find("meta", attrs={"name": name}) or soup.find("meta", attrs={"property": name})
        if m and m.get("content"):
            parts.append(m["content"])
    body = soup.body or soup
    parts.append(body.get_text(" ", strip=True))
    text = re.sub(r"\s+", " ", " ".join(parts)).strip()
    return text[:MAX_TEXT_CHARS]


def reg_host(url: str) -> str:
    h = (urlparse(url).hostname or "").lower()
    return h[4:] if h.startswith("www.") else h


def pick_internal_links(soup: BeautifulSoup, base_url: str) -> list[str]:
    base = reg_host(base_url)
    scored: dict[str, float] = {}
    nav_links = set()
    for container in soup.find_all(["nav", "header"]) + soup.select("[role=navigation], [class*=menu], [id*=menu]"):
        for a in container.find_all("a", href=True):
            nav_links.add(id(a))
    for a in soup.find_all("a", href=True):
        href = a["href"].strip()
        if not href or href.startswith(("mailto:", "tel:", "javascript:", "#")):
            continue
        url = urldefrag(urljoin(base_url, href))[0]
        p = urlparse(url)
        if p.scheme not in ("http", "https") or reg_host(url) != base:
            continue
        if SKIP_EXT.search(p.path) or p.path.rstrip("/") in ("", "/fr", "/en", "/fr-fr", "/en-us"):
            continue
        label = f"{p.path} {a.get_text(' ', strip=True)} {a.get('title', '')}"
        if EXCLUDE_LINK.search(label):
            continue
        s = sum(w for rx, w in WANTED_LINK if rx.search(label))
        if s == 0:
            continue
        if id(a) in nav_links:
            s += 2
        s -= 0.3 * p.path.strip("/").count("/")  # préfère les pages de premier niveau
        key = url.rstrip("/")
        scored[key] = max(scored.get(key, -1e9), s)
    best = sorted(scored.items(), key=lambda kv: -kv[1])
    return [u for u, _ in best[:MAX_INTERNAL_PAGES]]


# --------------------------------------------------------------------------- crawler
class Fetcher:
    def __init__(self, client: httpx.AsyncClient):
        self.client = client
        self.global_sem = asyncio.Semaphore(MAX_PARALLEL_REQUESTS)
        self.host_locks: dict[str, asyncio.Lock] = defaultdict(asyncio.Lock)
        self.host_last: dict[str, float] = {}
        self.robots: dict[str, RobotFileParser | None] = {}
        self.robots_locks: dict[str, asyncio.Lock] = defaultdict(asyncio.Lock)

    async def _request(self, url: str) -> httpx.Response:
        """Une seule requête HTTP, avec ≥1 s entre deux requêtes vers le même hôte et ≤10 en parallèle."""
        host = (urlparse(url).hostname or "").lower()
        async with self.host_locks[host]:
            wait = self.host_last.get(host, 0) + MIN_DELAY_PER_DOMAIN - time.monotonic()
            if wait > 0:
                await asyncio.sleep(wait)
            async with self.global_sem:
                self.host_last[host] = time.monotonic()
                try:
                    async with self.client.stream("GET", url) as r:
                        chunks, size = [], 0
                        ctype = r.headers.get("content-type", "")
                        if r.status_code < 300 and any(t in ctype for t in ("html", "xml", "text/plain")):
                            async for chunk in r.aiter_bytes():
                                chunks.append(chunk)
                                size += len(chunk)
                                if size > MAX_BYTES:
                                    break
                        r._content = b"".join(chunks)
                        return r
                finally:
                    self.host_last[host] = time.monotonic()

    async def get(self, url: str) -> httpx.Response:
        """GET avec redirections suivies manuellement (pour appliquer la limite par hôte à chaque saut)."""
        for _ in range(MAX_REDIRECTS + 1):
            r = await self._request(url)
            if r.is_redirect and r.headers.get("location"):
                url = urljoin(str(r.url), r.headers["location"])
                continue
            return r
        raise httpx.TooManyRedirects(f"plus de {MAX_REDIRECTS} redirections", request=r.request)

    async def allowed(self, url: str) -> bool:
        p = urlparse(url)
        origin = f"{p.scheme}://{p.netloc}"
        async with self.robots_locks[origin]:
            if origin not in self.robots:
                rp: RobotFileParser | None = RobotFileParser()
                try:
                    r = await self.get(origin + "/robots.txt")
                    if r.status_code in (401, 403):
                        rp.disallow_all = True
                    elif r.status_code >= 400 or not r.content:
                        rp.allow_all = True
                    else:
                        rp.parse(r.content.decode(r.encoding or "utf-8", "replace").splitlines())
                except Exception:
                    rp = None  # robots.txt inaccessible : on considère le site autorisé
                self.robots[origin] = rp
        rp = self.robots[origin]
        return True if rp is None else rp.can_fetch(USER_AGENT, url)

    async def fetch_page(self, url: str) -> tuple[str, BeautifulSoup | None, str | None]:
        """Retourne (url finale, soup, erreur)."""
        if not await self.allowed(url):
            return url, None, "interdit par robots.txt"
        r = await self.get(url)
        final = str(r.url)
        if r.status_code >= 400:
            return final, None, f"HTTP {r.status_code}"
        if "html" not in r.headers.get("content-type", "html"):
            return final, None, f"contenu non HTML ({r.headers.get('content-type')})"
        html = r.content.decode(r.encoding or "utf-8", "replace") if r.content else ""
        if not html:
            return final, None, "page vide"
        return final, BeautifulSoup(html, "lxml"), None


def normalize_site(url: str) -> str | None:
    if not isinstance(url, str) or not url.strip():
        return None
    url = url.strip()
    if not re.match(r"https?://", url, re.I):
        url = "http://" + url
    p = urlparse(url)
    if not p.hostname:
        return None
    return f"{p.scheme.lower()}://{p.netloc.lower()}{p.path or '/'}"


async def crawl_site(fetcher: Fetcher, site: str) -> dict:
    out = {"site": site, "final_url": None, "pages": [], "error": None}
    try:
        final, soup, err = await fetcher.fetch_page(site)
        if soup is None and site.startswith("http://") and err and not err.startswith("interdit"):
            # certains sites ne répondent qu'en https
            final2, soup2, err2 = await fetcher.fetch_page("https://" + site[len("http://"):])
            if soup2 is not None:
                final, soup, err = final2, soup2, None
        out["final_url"] = final
        if soup is None:
            out["error"] = err
            return out
        links = pick_internal_links(soup, final)
        out["pages"].append({"url": final, "text": extract_text(soup)})
        for link in links:
            try:
                f, s, e = await fetcher.fetch_page(link)
                if s is not None:
                    out["pages"].append({"url": f, "text": extract_text(s)})
                else:
                    out["pages"].append({"url": f, "text": "", "error": e})
            except Exception as e:  # une page interne en échec n'invalide pas le site
                out["pages"].append({"url": link, "text": "", "error": f"{type(e).__name__}: {e}"[:200]})
    except httpx.TimeoutException:
        out["error"] = "timeout"
    except httpx.ConnectError as e:
        out["error"] = f"connexion impossible: {e}"[:200]
    except Exception as e:
        out["error"] = f"{type(e).__name__}: {e}"[:200]
    return out


# --------------------------------------------------------------------------- cache SQLite
class Cache:
    def __init__(self, path: Path):
        self.db = sqlite3.connect(path)
        self.db.execute("CREATE TABLE IF NOT EXISTS sites (site TEXT PRIMARY KEY, data TEXT, fetched_at REAL)")
        self.db.commit()

    def get(self, site: str) -> dict | None:
        row = self.db.execute("SELECT data FROM sites WHERE site=?", (site,)).fetchone()
        return json.loads(row[0]) if row else None

    def put(self, site: str, data: dict) -> None:
        self.db.execute("INSERT OR REPLACE INTO sites VALUES (?,?,?)", (site, json.dumps(data, ensure_ascii=False),
                                                                        time.time()))
        self.db.commit()


async def crawl_all(sites: list[str], cache: Cache) -> None:
    limits = httpx.Limits(max_connections=MAX_PARALLEL_REQUESTS * 2, max_keepalive_connections=MAX_PARALLEL_REQUESTS)
    async with httpx.AsyncClient(headers=HEADERS, timeout=TIMEOUT, follow_redirects=False, limits=limits,
                                 http2=False) as client:
        fetcher = Fetcher(client)
        site_sem = asyncio.Semaphore(MAX_PARALLEL_REQUESTS * 3)

        async def one(site: str):
            async with site_sem:
                data = await crawl_site(fetcher, site)
            cache.put(site, data)
            pbar.update(1)
            if data["error"]:
                pbar.set_postfix_str(f"erreur {site[:40]}: {data['error'][:40]}")

        with tqdm(total=len(sites), desc="Sites", unit="site") as pbar:
            await asyncio.gather(*(one(s) for s in sites))


# --------------------------------------------------------------------------- résultats
def classify(data: dict | None) -> dict:
    if data is None:
        return {"IA": "indéterminé", "Data": "indéterminé", "erreur": "non traité"}
    pages = [p for p in data.get("pages", []) if p.get("text")]
    text = " \n ".join(p["text"] for p in pages)
    n_words = len(text.split())
    res = {"pages": " | ".join(p["url"] for p in pages), "nb_mots": n_words, "url_finale": data.get("final_url"),
           "erreur": data.get("error") or ""}
    if data.get("error") or n_words < MIN_WORDS:
        if not res["erreur"]:
            res["erreur"] = f"texte insuffisant ({n_words} mots, site probablement en JavaScript)"
        res.update({"IA": "indéterminé", "Data": "indéterminé"})
        return res
    ia, dt = score_text(text, IA_COMPILED), score_text(text, DATA_COMPILED)
    res.update({
        "IA": decide(ia), "score_IA": ia["score"], "mots_cles_IA": ia["keywords"],
        "Data": decide(dt), "score_Data": dt["score"], "mots_cles_Data": dt["keywords"],
    })
    return res


OUT_COLUMNS = [
    ("IA", "IA"), ("score_IA", "Score IA"), ("mots_cles_IA", "Mots-clés IA"),
    ("Data", "Data"), ("score_Data", "Score Data"), ("mots_cles_Data", "Mots-clés Data"),
    ("IA_ou_Data", "IA ou Data"), ("url_finale", "URL finale"), ("pages", "Pages analysées"),
    ("nb_mots", "Nb mots"), ("erreur", "Erreur"),
]


def combine(ia: str, data: str) -> str:
    order = ["oui", "probable", "non", "indéterminé"]
    return min((ia, data), key=order.index)


def write_excel(df: pd.DataFrame, path: Path) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    df.to_excel(path, index=False, sheet_name="Résultats")
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb.active
    fills = {"oui": "C6EFCE", "probable": "FFEB9C", "non": "F2F2F2", "indéterminé": "FFC7CE"}
    status_cols = [i + 1 for i, c in enumerate(df.columns) if c in ("IA", "Data", "IA ou Data")]
    for c in range(1, ws.max_column + 1):
        ws.cell(1, c).font = Font(bold=True)
        width = 14 if df.columns[c - 1] not in ("Mots-clés IA", "Mots-clés Data", "Pages analysées", "Erreur",
                                                "RAISON SOCIALE", "CLASSIFICATION") else 45
        ws.column_dimensions[get_column_letter(c)].width = width
    for r in range(2, ws.max_row + 1):
        for c in status_cols:
            v = ws.cell(r, c).value
            if v in fills:
                ws.cell(r, c).fill = PatternFill("solid", fgColor=fills[v])
                ws.cell(r, c).alignment = Alignment(horizontal="center")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def choose_sample(df: pd.DataFrame, n: int) -> pd.DataFrame:
    """Échantillon varié : réparti sur les tranches d'effectif, avec quelques noms évoquant l'IA/data."""
    hint = df["RAISON SOCIALE"].str.contains(r"\b(?:data|ia|ai|intelligence|analytics|cognitive|quant)",
                                             case=False, regex=True, na=False)
    picked = df[hint].sample(min(n // 4, hint.sum()), random_state=42)
    rest = df.drop(picked.index)
    col = "TRANCHE D’EFFECTIF ENTREPRISE" if "TRANCHE D’EFFECTIF ENTREPRISE" in df.columns else None
    k = n - len(picked)
    if col:
        idx = []
        for _, g in rest.groupby(col):
            idx += g.sample(min(len(g), max(1, round(k * len(g) / len(rest)))), random_state=42).index.tolist()
        others = rest.loc[idx].head(k)
    else:
        others = rest.sample(k, random_state=42)
    return pd.concat([picked, others]).sort_index()


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input", type=Path)
    ap.add_argument("-o", "--output", type=Path, help="fichier Excel de sortie (défaut : <input>_ia_data.xlsx)")
    ap.add_argument("--cache", type=Path, help="cache SQLite (défaut : <input>_cache.sqlite)")
    ap.add_argument("--sample", type=int, help="ne traiter qu'un échantillon varié de N entreprises")
    ap.add_argument("--limit", type=int, help="ne traiter que les N premières entreprises")
    ap.add_argument("--rescore-only", action="store_true", help="pas de téléchargement, recalcule depuis le cache")
    ap.add_argument("--retry-errors", action="store_true", help="re-télécharge les sites en erreur dans le cache")
    ap.add_argument("--site-col", default="SITE WEB")
    ap.add_argument("--filter-col", default="categorie", help="colonne de filtre (insensible à la casse)")
    ap.add_argument("--filter-values", default="OK,A_VERIFIER")
    args = ap.parse_args()

    output = args.output or args.input.with_name(args.input.stem + "_ia_data.xlsx")
    cache_path = args.cache or args.input.with_name(args.input.stem + "_cache.sqlite")
    if output.resolve() == args.input.resolve():
        sys.exit("Le fichier de sortie doit être différent du fichier d'origine.")

    df = pd.read_excel(args.input)  # lecture seule : le fichier d'origine n'est jamais réécrit
    fcol = next((c for c in df.columns if str(c).strip().lower() == args.filter_col.lower()), None)
    if fcol:
        keep = {v.strip().upper() for v in args.filter_values.split(",")}
        df = df[df[fcol].astype(str).str.strip().str.upper().isin(keep)]
        print(f"Filtre {fcol} ∈ {sorted(keep)} : {len(df)} entreprises")
    else:
        print(f"Pas de colonne « {args.filter_col} » : toutes les {len(df)} entreprises sont traitées")
    if args.sample:
        df = choose_sample(df, args.sample)
    elif args.limit:
        df = df.head(args.limit)

    df = df.copy()
    df["_site"] = df[args.site_col].map(normalize_site)
    cache = Cache(cache_path)
    sites = sorted({s for s in df["_site"].dropna()})
    if not args.rescore_only:
        todo = []
        for s in sites:
            c = cache.get(s)
            if c is None or (args.retry_errors and c.get("error")):
                todo.append(s)
        print(f"{len(sites)} sites uniques, {len(sites) - len(todo)} déjà en cache, {len(todo)} à télécharger")
        if todo:
            try:
                asyncio.run(crawl_all(todo, cache))
            except KeyboardInterrupt:
                print("\nInterrompu : les sites déjà traités sont dans le cache, relancez pour reprendre.")

    results = [classify(cache.get(s)) if s else {"IA": "indéterminé", "Data": "indéterminé", "erreur": "pas de site"}
               for s in df["_site"]]
    res = pd.DataFrame(results, index=df.index)
    res["IA_ou_Data"] = [combine(a, b) for a, b in zip(res["IA"], res["Data"])]
    out = df.drop(columns="_site").copy()
    for key, name in OUT_COLUMNS:
        out[name] = res[key] if key in res else None
    write_excel(out, output)
    print(f"\nRésultats écrits dans {output}")
    for col in ("IA", "Data", "IA ou Data"):
        print(f"  {col:10s}", out[col].value_counts().to_dict())


if __name__ == "__main__":
    main()
